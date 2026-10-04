"""Workbook templates: make one from a finished workbook, apply it to new data.

Not part of the R package. Modelled on Tableau's Accelerators (a list of
required fields, a type-checked mapping, unmapped fields break the sheets
that use them), Power BI `.pbit` files (no data inside, parameter values
asked for on use), Grafana's "export for sharing" (connections become
placeholders) and Copier (the answers are saved in the output so the
template can be re-applied later).

A template is an ordinary `.twbx` (Tableau still opens it) with a
`template.json` manifest inside. Applying it keeps every sheet, dashboard,
calculation and format of the template and points its fields at the new
data: each field keeps the local name that sheets and formulas use, and only
the physical column behind it (the metadata record's `remote-name`) and the
connection change. This mirrors what Tableau's Replace Data Source does, but
the output has not been opened in Tableau itself.
"""

from __future__ import annotations

import copy
import datetime as _dt
import difflib
import hashlib
import io
import json
import os
import re
import uuid
import warnings
import zipfile
from dataclasses import dataclass, field as _field
from pathlib import Path
from typing import Iterable, Optional, Union

import pandas as pd
from lxml import etree

from ._clean import is_missing
from .parser import TwbParser
from .rename import _match_key, _words
from .usage import field_usage

MANIFEST_NAME = "template.json"
ANSWERS_NAME = "template-answers.json"
TEMPLATE_FORMAT = "py-tbparse-template"
MANIFEST_VERSION = 2   # 2 adds template id/revision and a uid per field; version 1 still loads

MAPPING_COLUMNS = [
    "datasource", "field", "caption", "datatype", "required", "used_by",
    "mapped_to", "data_type", "status", "score",
]

_AUTO_COLUMN = "{http://www.tableausoftware.com/xml/user}auto-column"
_SECRET_ATTRS = ("username", "password")
_CONNECTION_ATTRS = ("class", "server", "dbname", "schema", "port", "directory", "filename",
                     "warehouse", "service", "authentication")
# Members of a .twbx that hold data rather than the workbook's own assets.
_DATA_SUFFIXES = (".hyper", ".tde", ".csv", ".txt", ".tsv", ".xlsx", ".xls", ".xlsm", ".json",
                  ".zip", ".shp", ".shx", ".dbf", ".prj", ".kml", ".geojson", ".mdb", ".accdb", ".sav")

# Tableau's remote-type codes for the file connections we write. In the 200-workbook corpus a text
# file's date column is 133 (62 records, 28 workbooks) far more often than 7 (7 records, 5 workbooks).
_REMOTE_TYPE = {"string": 129, "integer": 20, "real": 5, "boolean": 11, "date": 133, "datetime": 135}
# The Excel driver's, measured over the 88 `excel-direct` workbooks of the corpus (1,770 columns): the code, and
# the `DebugRemoteType` attribute Tableau writes with it. Dates and date-times are both DATE (7).
_EXCEL_REMOTE_TYPE = {"string": (130, "WSTR"), "integer": (20, "I8"), "real": (5, "R8"), "boolean": (11, "WINBOOL"),
                      "date": (7, "DATE"), "datetime": (7, "DATE")}
_NUMERIC = {"integer", "real"}
_DATES = {"date", "datetime"}


class TemplateError(ValueError):
    """The template cannot be made or applied as asked."""


# ---------------------------------------------------------------- helpers --

def _now() -> str:
    return _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0).isoformat()


def _version() -> str:
    from . import __version__

    return __version__


def _new_id() -> str:
    return str(uuid.uuid4())


def field_uid(datasource: str, name: str, role: str, datatype: str) -> str:
    """A field's identity across template revisions: it survives a caption change, not a change of
    the field's local name, role or type."""
    raw = "\x1f".join((datasource or "", name or "", role or "", datatype or ""))
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def _endswith_tag(el, name: str) -> bool:
    # Tableau prefixes some tags with feature flags: `_.fcp.Flag.true...object-graph`.
    return isinstance(el.tag, str) and (el.tag == name or el.tag.endswith("..." + name))


def _text_of(el) -> str:
    return " ".join(t.strip() for t in el.itertext() if t.strip()) if el is not None else ""


def _role(datatype: Optional[str]) -> str:
    return "measure" if datatype in _NUMERIC else "dimension"


def _non_parameter_datasources(doc) -> list:
    return [ds for ds in doc.xpath("/workbook/datasources/datasource[@name]") if ds.get("name") != "Parameters"]


def _physical_fields(ds) -> list[dict]:
    """The fields a datasource takes straight from its data, from the
    metadata records that link each physical column to a local name."""
    out, seen = [], set()
    for rec in ds.xpath("./connection//metadata-record[@class='column']"):
        local = rec.findtext("local-name")
        if not local or local in seen:
            continue
        seen.add(local)
        out.append({
            "name": local,
            "remote": rec.findtext("remote-name") or local.strip("[]"),
            "table": rec.findtext("parent-name"),
            "datatype": rec.findtext("local-type") or "string",
            "alias": rec.findtext("remote-alias"),
        })
    return out


# ------------------------------------------------------------------ make --

def _parameters(doc) -> list[dict]:
    out = []
    for col in doc.xpath("/workbook/datasources/datasource[@name='Parameters']/column[@name]"):
        members = [m.get("value") for m in col.xpath("./members/member[@value]")]
        rng = col.find("range")
        out.append({
            "name": col.get("name"),
            "caption": col.get("caption") or col.get("name").strip("[]"),
            "datatype": col.get("datatype"),
            "value": col.get("value"),
            "allowed": members or None,
            "range": {k: rng.get(k) for k in ("min", "max", "granularity") if rng.get(k) is not None}
            if rng is not None else None,
            "description": _text_of(col.find("desc")) or None,
        })
    return out


def _connections(ds) -> list[dict]:
    """Where a datasource's data lives, without credentials."""
    conns = ds.xpath("./connection/named-connections/named-connection/connection") or ds.xpath("./connection")
    out = []
    for c in conns:
        info = {k: c.get(k) for k in _CONNECTION_ATTRS if c.get(k)}
        if info.get("class") != "federated":
            out.append(info)
    return out


def _scrub(doc) -> None:
    for el in doc.iter():
        if isinstance(el.tag, str) and el.tag == "connection":
            for attr in _SECRET_ATTRS:
                if el.get(attr):
                    el.set(attr, "")


def _strip_extracts(doc) -> int:
    n = 0
    for ds in doc.xpath("/workbook/datasources/datasource"):
        for el in list(ds):
            if _endswith_tag(el, "extract"):
                ds.remove(el)
                n += 1
    return n


def build_manifest(parser: TwbParser, name: Optional[str] = None, description: Optional[str] = None,
                   template_id: Optional[str] = None, revision: int = 1) -> dict:
    """The manifest `make_template` writes: what a template needs to be applied."""
    doc = parser.xml_doc
    usage = field_usage(parser)
    used = {(r.datasource, r.field): r for r in usage.itertuples(index=False)}
    datasources = []
    for ds in _non_parameter_datasources(doc):
        dsname = ds.get("name")
        cols = {c.get("name"): c for c in ds.xpath("./column[@name]")}
        fields = []
        for f in _physical_fields(ds):
            col = cols.get(f["name"])
            if col is not None and col.get("{http://www.tableausoftware.com/xml/user}auto-column"):
                continue  # Tableau's own [Number of Records] etc.; apply gives it its formula back
            u = used.get((dsname, f["name"]))
            datatype = (col.get("datatype") if col is not None else None) or f["datatype"]
            role = (col.get("role") if col is not None else None) or _role(f["datatype"])
            fields.append({
                "name": f["name"],
                "uid": field_uid(dsname, f["name"], role, datatype),
                "caption": col.get("caption") if col is not None else None,
                "remote": f["remote"],
                "alias": f.get("alias"),
                "datatype": datatype,
                # what the data delivered; differs from `datatype` when the author retyped the field
                "physical_type": f["datatype"],
                "customized": bool(col is not None and col.get("datatype-customized") == "true"),
                "role": role,
                "required": bool(u and u.used),
                "used_by": list(u.sheets) if u else [],
                "description": _text_of(col.find("desc")) if col is not None and col.find("desc") is not None else None,
            })
        datasources.append({
            "name": dsname,
            "caption": ds.get("caption"),
            "connections": _connections(ds),
            "fields": fields,
        })
    src = Path(parser.twbx_path or parser.path)
    return {
        "format": TEMPLATE_FORMAT,
        "version": MANIFEST_VERSION,
        "id": template_id or _new_id(),
        "revision": revision,
        "name": name or src.stem,
        "description": description or "",
        "source": src.name,
        "created": _now(),
        "created_with": f"py-tbparse {_version()}",
        "datasources": datasources,
        "parameters": _parameters(doc),
        "worksheets": doc.xpath("/workbook/worksheets/worksheet/@name"),
        "dashboards": doc.xpath("/workbook/dashboards/dashboard/@name"),
    }


def _is_data_member(name: str) -> bool:
    low = name.lower()
    # TwbxExternalCache holds cached query results: data too.
    return low.startswith(("data/", "twbxexternalcache/")) or low.endswith(_DATA_SUFFIXES)


def _write_new(path: Path, data: bytes, overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(f"refusing to overwrite existing file: {path}")
    with open(path, "wb" if overwrite else "xb") as fh:
        fh.write(data)


def _package(parser: TwbParser, twb: bytes, extra: dict[str, bytes], keep_data: bool,
             drop: Iterable[str] = ()) -> bytes:
    """A .twbx holding `twb` plus `extra` members; the source's other
    members are copied (data files only with `keep_data`)."""
    out = io.BytesIO()
    drop = set(drop) | set(extra)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        if parser.twbx_path:
            twb_name = parser.twb_name
            with zipfile.ZipFile(parser.twbx_path) as src:
                for info in src.infolist():
                    if info.filename == twb_name:
                        dst.writestr(info, twb, compress_type=info.compress_type)
                    elif info.filename in drop or (not keep_data and _is_data_member(info.filename)):
                        continue
                    else:
                        dst.writestr(info, src.read(info.filename), compress_type=info.compress_type)
        else:
            dst.writestr(_member(Path(parser.path).name), twb)
        for name, data in extra.items():
            dst.writestr(_member(name), data)
    return out.getvalue()


def _member(name: str) -> zipfile.ZipInfo:
    """A zip entry with a fixed timestamp, so the same inputs always make the same bytes."""
    info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    return info


def default_template_path(parser: TwbParser) -> str:
    src = Path(parser.twbx_path or parser.path)
    return str(src.with_name(f"{src.stem}.template.twbx"))


def make_template(
    parser: Union[TwbParser, str],
    output_path: Optional[str] = None,
    name: Optional[str] = None,
    description: Optional[str] = None,
    keep_data: bool = False,
    overwrite: bool = False,
    template_id: Optional[str] = None,
    revision_of: Union["Template", str, None] = None,
) -> str:
    """Save a template made from a finished workbook; return its path.

    The template is a `.twbx` (default `<name>.template.twbx` beside the
    source) holding a copy of the workbook and a `template.json` manifest
    that lists the fields the workbook needs (`required` when a sheet uses
    them, directly or through calculations, groups and sets), its
    parameters and where its data came from. Passwords and user names are
    blanked. Extracts and packaged data files are left out unless
    `keep_data` (handy for an Accelerator-style template that opens with
    sample data). The source is never modified and nothing is overwritten
    unless `overwrite=True`. The manifest carries an `id` (`template_id`, or
    a new UUID) that stays the same across revisions of one template, and a
    `revision` number. `revision_of` (a template or its path) makes this the
    next revision of that template: same `id`, `revision` one higher, so
    workbooks made from the old one can be brought up to date
    (`template_update.update_from_answers`). The old template has no `id`
    when it came from 0.4.x; the new one then gets a fresh one.
    """
    revision = 1
    if revision_of is not None:
        previous = revision_of if isinstance(revision_of, Template) else load_template(str(revision_of))
        if template_id and previous.id and template_id != previous.id:
            raise TemplateError(f"template_id {template_id!r} contradicts revision_of (id {previous.id!r})")
        template_id = template_id or previous.id
        revision = previous.revision + 1
    if not isinstance(parser, TwbParser):
        parser = TwbParser(str(parser))
    out = Path(output_path) if output_path else Path(default_template_path(parser))
    if out.suffix.lower() != ".twbx":
        raise TemplateError(f"a template is a .twbx file, got {out.suffix or 'no extension'}")
    if out.resolve() == Path(parser.twbx_path or parser.path).resolve():
        raise FileExistsError(f"refusing to overwrite the source workbook: {out}")
    manifest = build_manifest(parser, name=name, description=description, template_id=template_id,
                              revision=revision)
    doc = copy.deepcopy(parser.xml_doc)
    _scrub(doc)
    if not keep_data:
        _strip_extracts(doc)
    twb = etree.tostring(doc, xml_declaration=True, encoding="utf-8")
    data = _package(parser, twb, {MANIFEST_NAME: json.dumps(manifest, indent=2).encode("utf-8")},
                    keep_data=keep_data, drop=[ANSWERS_NAME])
    _write_new(out, data, overwrite)
    return str(out)


# ------------------------------------------------------------------ load --

@dataclass
class Template:
    path: str
    parser: TwbParser
    manifest: dict
    manifest_sha256: str = ""
    _cache: dict = _field(default_factory=dict, repr=False, compare=False)

    @property
    def name(self) -> str:
        return self.manifest.get("name") or Path(self.path).stem

    @property
    def id(self) -> Optional[str]:
        """Identity of the template across revisions; None for a version 1 manifest, which had none."""
        return self.manifest.get("id")

    @property
    def revision(self) -> int:
        return int(self.manifest.get("revision") or 1)

    def datasource(self, which: Optional[str] = None) -> dict:
        """The manifest entry of one template datasource, by internal name or
        caption; with no name, the one that has required fields (or the
        only one)."""
        entries = self.manifest.get("datasources", [])
        if which:
            for e in entries:
                if which in (e["name"], e.get("caption")):
                    return e
            raise TemplateError(f"template has no datasource {which!r}; it has: "
                                + ", ".join(e.get("caption") or e["name"] for e in entries))
        with_required = [e for e in entries if any(f["required"] for f in e["fields"])]
        pick = with_required or entries
        if len(pick) != 1:
            raise TemplateError("template has several datasources; pick one with datasource=: "
                                + ", ".join(e.get("caption") or e["name"] for e in pick))
        return pick[0]

    def fields(self) -> pd.DataFrame:
        rows = []
        for ds in self.manifest.get("datasources", []):
            for f in ds["fields"]:
                rows.append({"datasource": ds.get("caption") or ds["name"], "field": f["name"],
                             "caption": f.get("caption"), "datatype": f["datatype"],
                             "required": f["required"], "used_by": "; ".join(f.get("used_by") or [])})
        return pd.DataFrame(rows, columns=["datasource", "field", "caption", "datatype", "required", "used_by"])

    def parameters(self) -> pd.DataFrame:
        rows = [{"parameter": p["caption"], "datatype": p["datatype"], "value": p["value"],
                 "allowed": "; ".join(p["allowed"]) if p.get("allowed") else ""}
                for p in self.manifest.get("parameters", [])]
        return pd.DataFrame(rows, columns=["parameter", "datatype", "value", "allowed"])


def _fill_version_1(manifest: dict) -> None:
    """What a version 1 manifest lacks, worked out the way `make` would have: a uid per field and a
    revision. The `id` stays absent (None): there is nothing to derive a stable one from. The manifest's
    `version` is left as read, so the file on disk is never rewritten."""
    manifest.setdefault("revision", 1)
    for ds in manifest.get("datasources", []):
        for f in ds.get("fields", []):
            f.setdefault("uid", field_uid(ds["name"], f["name"], f.get("role"), f.get("datatype")))
            f.setdefault("alias", None)


def load_template(path: str) -> Template:
    """Open a template made by `make_template`."""
    if not str(path).lower().endswith(".twbx"):
        raise TemplateError(f"not a template (.twbx expected): {path}")
    with zipfile.ZipFile(path) as z:
        if MANIFEST_NAME not in z.namelist():
            raise TemplateError(f"{path} has no {MANIFEST_NAME}; make one with `py-tbparse template make`")
        raw = z.read(MANIFEST_NAME)
    manifest = json.loads(raw.decode("utf-8"))
    if manifest.get("format") != TEMPLATE_FORMAT:
        raise TemplateError(f"{path}: unknown template format {manifest.get('format')!r}")
    if int(manifest.get("version", 0)) > MANIFEST_VERSION:
        raise TemplateError(f"{path} was made by a newer py-tbparse (manifest v{manifest['version']})")
    _fill_version_1(manifest)
    return Template(path=str(path), parser=TwbParser(str(path)), manifest=manifest,
                    manifest_sha256=hashlib.sha256(raw).hexdigest())


# ------------------------------------------------------------- the data --

@dataclass
class DataSource:
    """New data for a template: its fields, and how to connect to it."""

    path: str
    kind: str  # "csv", "excel" or "tableau"
    fields: list[dict] = _field(default_factory=list)  # {name, datatype}
    element: object = None  # the Tableau <datasource> for kind "tableau"
    sheet: Optional[str] = None  # the worksheet, for kind "excel"
    grid: Optional[str] = None   # where the header and data sit on that sheet, as Excel writes it: "A1:D11"

    def names(self) -> list[str]:
        return [f["name"] for f in self.fields]

    def column(self, name: str) -> str:
        """The data's column a mapping means by `name`. `load_mapping` trims the spaces around what a person
        typed, so a header such as `'Date of Birth '` comes back as `'Date of Birth'`: take the exact column
        if there is one, else the only column that trims to `name`, else `name` itself."""
        names = self.names()
        if name in names:
            return name
        trimmed = [n for n in names if n.strip() == name]
        return trimmed[0] if len(trimmed) == 1 else name

    def fingerprint(self) -> str:
        """A hash of the column names and types: when it changes between two runs, the data's shape did."""
        raw = "\x1e".join(f"{f['name']}\x1f{f['datatype'] or ''}" for f in self.fields)
        return "sha256:" + hashlib.sha256(raw.encode("utf-8")).hexdigest()

    def role(self, name: str) -> Optional[str]:
        """`dimension` or `measure` when the data says so (a Tableau data source does; a CSV does not)."""
        for f in self.fields:
            if f["name"] == name:
                return f.get("role")
        return None

    def datatype(self, name: str) -> Optional[str]:
        for f in self.fields:
            if f["name"] == name:
                return f["datatype"]
        return None


def _infer_csv_type(series: pd.Series) -> Optional[str]:
    """Tableau's type for a CSV column; None when it has no values to judge by."""
    if series.dropna().astype(str).str.strip().eq("").all():
        return None
    if pd.api.types.is_bool_dtype(series):
        return "boolean"
    if pd.api.types.is_integer_dtype(series):
        return "integer"
    if pd.api.types.is_float_dtype(series):
        # whole numbers with blanks come back as float
        s = series.dropna()
        return "integer" if len(s) and (s == s.round()).all() and series.isna().any() else "real"
    s = series.dropna().astype(str)
    if s.str.lower().isin(["true", "false"]).all():
        return "boolean"
    if s.str.match(r"^\d{4}-\d{1,2}-\d{1,2}([ T]\d{1,2}:\d{2}(:\d{2})?)?$").all() or \
            s.str.match(r"^\d{1,2}/\d{1,2}/\d{2,4}( \d{1,2}:\d{2}(:\d{2})?)?$").all():
        return "datetime" if s.str.contains(":").any() else "date"
    return "string"


def _excel_type(values: list) -> Optional[str]:
    """Tableau's type for an Excel column, from the Python values openpyxl read; None when it has none."""
    vals = [v for v in values if v is not None and not (isinstance(v, str) and not v.strip())]
    if not vals:
        return None
    if all(isinstance(v, bool) for v in vals):
        return "boolean"
    if any(isinstance(v, bool) for v in vals):
        return "string"
    if all(isinstance(v, (_dt.datetime, _dt.date)) for v in vals):
        timed = any(isinstance(v, _dt.datetime) and (v.hour, v.minute, v.second) != (0, 0, 0) for v in vals)
        return "datetime" if timed else "date"
    if all(isinstance(v, (int, float)) for v in vals):
        # Excel keeps one kind of number; a column of whole numbers is an integer to Tableau
        return "integer" if all(isinstance(v, int) or float(v).is_integer() for v in vals) else "real"
    return "string"


def _read_excel(p: Path, sheet: Union[str, int, None]) -> DataSource:
    """The columns of one worksheet of an `.xlsx`/`.xlsm` file: the first non-empty row is the header, the
    next `_SAMPLE_ROWS` rows decide each column's type."""
    if p.suffix.lower() in (".xls", ".xlsb"):
        raise TemplateError(f"{p.name}: old Excel formats ({p.suffix}) are not read; save it as .xlsx or .csv")
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise TemplateError(f"reading {p.name} needs openpyxl: pip install 'py-tbparse[excel]'") from None
    try:
        book = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, openpyxl.utils.exceptions.InvalidFileException) as e:
        raise TemplateError(f"{p.name} is not a readable Excel file: {e}") from None
    try:
        sheets = {ws.title: ws for ws in book.worksheets}
        names = list(sheets)
        if sheet is None:
            visible = [n for n in names if sheets[n].sheet_state == "visible"]
            if len(visible) != 1:
                raise TemplateError(f"{p.name}: " + (
                    f"{len(visible)} visible sheets ({', '.join(visible)}); pick one with sheet=" if visible
                    else "no visible sheet; name one with sheet="))
            chosen = visible[0]
        elif isinstance(sheet, int):
            if not -len(names) <= sheet < len(names):
                raise TemplateError(f"{p.name} has {len(names)} sheet(s), none at index {sheet}")
            chosen = names[sheet]
        elif sheet in sheets:
            chosen = sheet
        else:
            raise TemplateError(f"{p.name} has no sheet {sheet!r}; it has: {', '.join(names)}")
        ws = sheets[chosen]
        header = None
        rows: list = []
        last = 0
        more = False
        for number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            filled = any(v is not None and (not isinstance(v, str) or v.strip()) for v in row)
            if header is None:
                if filled:
                    header = (number, row)
                continue
            if len(rows) >= _SAMPLE_ROWS:
                more = True
                break
            rows.append(row)
            if filled:
                last = number
        if header is None:
            raise TemplateError(f"{p.name}: sheet {chosen!r} is empty")
        top, cells = header
        used = [i for i, v in enumerate(cells) if v is not None and (not isinstance(v, str) or v.strip())]
        first, end = used[0], used[-1] + 1
        names_out: list[str] = []
        for i in range(first, end):
            v = cells[i]
            blank = v is None or not str(v).strip()
            name = f"F{i - first + 1}" if blank else str(v)
            base, k = name, 0
            while name in names_out:
                k += 1
                name = f"{base}{k}"
            if name != base and not blank:
                warnings.warn(f"{p.name}: sheet {chosen!r} repeats the header {base!r}; the later column is {name!r}")
            names_out.append(name)
        fields = [{"name": name, "datatype": _excel_type([r[first + j] if len(r) > first + j else None for r in rows])}
                  for j, name in enumerate(names_out)]
        bottom = (ws.max_row or last or top) if more else max(last, top)
        grid = f"{get_column_letter(first + 1)}{top}:{get_column_letter(end)}{bottom}"
        return DataSource(path=str(p.resolve()), kind="excel", fields=fields, sheet=chosen, grid=grid)
    finally:
        book.close()


def read_data(path: str, datasource: Optional[str] = None, sheet: Union[str, int, None] = None) -> DataSource:
    """Describe new data for `apply_template`: a CSV file, an Excel file
    (`.xlsx`/`.xlsm`, needs `py-tbparse[excel]`; `sheet=` picks a worksheet by
    name or by index from 0, and is required when several are visible), or a
    Tableau workbook (`.twb`/`.twbx`) or data source (`.tds`) already
    connected to it (pick one of several datasources with `datasource=`)."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"no such file: {path}")
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xls", ".xlsb"):
        return _read_excel(p, sheet)
    if suffix in (".csv", ".txt", ".tsv"):
        sep = "\t" if suffix == ".tsv" else ","
        df = pd.read_csv(p, sep=sep, nrows=2000, encoding="utf-8-sig")
        return DataSource(path=str(p.resolve()), kind="csv",
                          fields=[{"name": str(c), "datatype": _infer_csv_type(df[c])} for c in df.columns])
    if suffix in (".twb", ".twbx"):
        doc = TwbParser(str(p)).xml_doc
        candidates = _non_parameter_datasources(doc)
    elif suffix == ".tds":
        root = etree.parse(str(p)).getroot()
        candidates = [root] if root.tag == "datasource" else root.xpath("//datasource[@name]")
    else:
        raise TemplateError(f"unsupported data file {p.name}: use a .csv, .xlsx, .twb, .twbx or .tds")
    if datasource:
        candidates = [d for d in candidates if datasource in (d.get("name"), d.get("caption"), d.get("formatted-name"))]
    candidates = [d for d in candidates if d.find("connection") is not None]
    if len(candidates) != 1:
        names = [d.get("caption") or d.get("name") for d in candidates]
        raise TemplateError(f"{p.name}: expected one datasource with a connection, found {len(candidates)}"
                            + (f" ({', '.join(names)}); pick one" if names else ""))
    ds = candidates[0]
    roles = {c.get("name"): c.get("role") for c in ds.xpath("./column[@name][@role]")}
    fields = [{"name": f["remote"], "datatype": f["datatype"], "local": f["name"], "role": roles.get(f["name"])}
              for f in _physical_fields(ds)]
    if not fields:
        raise TemplateError(f"{p.name}: the datasource lists no columns (open it in Tableau once and save)")
    return DataSource(path=str(p.resolve()), kind="tableau", fields=fields, element=ds)


# --------------------------------------------------------------- mapping --

def _type_status(want: str, got: Optional[str]) -> Optional[str]:
    """None when compatible; otherwise why not (a "differs" status still maps)."""
    if got is None or want == got:
        return None
    if want in _NUMERIC and got in _NUMERIC:
        return "numeric type differs"
    if want in _DATES and got in _DATES:
        return "date type differs"
    return "type mismatch"


def suggest_mapping(
    template: Template,
    data: DataSource,
    datasource: Optional[str] = None,
    fuzzy_cutoff: float = 0.85,
) -> pd.DataFrame:
    """Pair each field the template needs with a column of the new data.

    A field matches a column whose name, ignoring case and separators,
    equals the field's caption, local name or original column name
    (`matched`); failing that, the closest name above `fuzzy_cutoff`
    (`close match`, never across a different number). Types are checked like
    Tableau's Accelerator mapper: a string column is never offered for a
    number or a date (`type mismatch`, left unmapped); integer vs real and
    date vs datetime map with a warning status, and so does a column that
    Tableau says is a dimension where the template has a measure, or the
    reverse (`role differs`; only a Tableau data source states a role, a
    CSV does not). The field's alias is tried after its caption, name and
    original column name. Each column is used once, required fields first;
    between equally close columns the one that comes first wins. Unmapped
    fields have `status` `missing` (required) or `unused` (not used by any
    sheet).

    Returns `datasource, field, caption, datatype, required, used_by,
    mapped_to, data_type, status, score`.
    """
    if not 0 <= fuzzy_cutoff <= 1:
        raise ValueError(f"fuzzy_cutoff must be between 0 and 1, got {fuzzy_cutoff}")
    entry = template.datasource(datasource)
    by_key: dict[str, str] = {}
    for n in data.names():
        by_key.setdefault(_match_key(n), n)
    keys = list(by_key)
    position = {n: i for i, n in enumerate(data.names())}   # ties go to the column that comes first
    taken: set = set()
    taken_by: dict = {}
    rows = []
    fields = sorted(entry["fields"], key=lambda f: not f["required"])
    for f in fields:
        names = [f.get("caption"), f["name"].strip("[]"), f.get("remote"), f.get("alias")]
        names = list(dict.fromkeys(n for n in names if n))
        pick, status, score = None, None, None
        for n in names:
            k = _match_key(n)
            if k in by_key and by_key[k] not in taken:
                pick, status, score = by_key[k], "matched", 1.0
                break
        if pick is None:
            best = None
            for n in names:
                k = _match_key(n)
                if not k:
                    continue
                digits = re.findall(r"\d+", k)
                for cand in difflib.get_close_matches(k, keys, n=5, cutoff=fuzzy_cutoff):
                    if re.findall(r"\d+", cand) != digits or by_key[cand] in taken:
                        continue
                    ratio = difflib.SequenceMatcher(None, k, cand).ratio()
                    rank = (ratio, -position[by_key[cand]])
                    if best is None or rank > best[1]:
                        best = (by_key[cand], rank)
            if best:
                pick, status, score = best[0], "close match", round(best[1][0], 3)
        if pick is not None:
            # Compare like with like: the column the template was built on vs the new one.
            # A field the author retyped keeps that type; Tableau converts the new column.
            got = data.datatype(pick)
            problems = [_type_status(f.get("physical_type") or f["datatype"], got)]
            if f.get("customized"):  # a column already of the field's own type also works
                problems.append(_type_status(f["datatype"], got))
            problem = None if None in problems else min(problems, key=lambda s: s == "type mismatch")
            if problem == "type mismatch":
                pick, status, score = None, f"type mismatch ({data.datatype(pick)} column {pick!r})", None
            elif problem:
                status = problem
            elif data.role(pick) and f.get("role") and data.role(pick) != f["role"]:
                status = "role differs"   # a dimension where the template has a measure, or the reverse; still mapped
        if pick is not None:
            taken.add(pick)
            taken_by[pick] = f["name"]
        elif status is None:
            status = "missing" if f["required"] else "unused"
            for n in names:  # the right column exists but another field already took it
                col = by_key.get(_match_key(n))
                if col in taken_by:
                    status = f"column {col!r} already used by {taken_by[col]}"
                    break
        rows.append({
            "datasource": entry["name"], "field": f["name"], "caption": f.get("caption"),
            "datatype": f["datatype"], "required": bool(f["required"]),
            "used_by": "; ".join(f.get("used_by") or []),
            "mapped_to": pick or "", "data_type": data.datatype(pick) if pick else "",
            "status": status, "score": score,
        })
    return pd.DataFrame(rows, columns=MAPPING_COLUMNS)


def load_mapping(source: Union[str, os.PathLike, pd.DataFrame]) -> pd.DataFrame:
    """Read a mapping back in (the CSV `suggest_mapping` / `template apply
    --mapping-out` wrote, with `mapped_to` edited by hand; blank means no
    column)."""
    df = source.copy() if isinstance(source, pd.DataFrame) else pd.read_csv(
        source, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in ("field", "mapped_to") if c not in df.columns]
    if missing:
        raise ValueError(f"mapping is missing column(s): {', '.join(missing)}")
    df["mapped_to"] = df["mapped_to"].map(lambda v: "" if is_missing(v) else str(v).strip())
    used = df[df["mapped_to"] != ""]["mapped_to"]
    dup = used[used.duplicated()]
    if len(dup):
        raise ValueError(f"column {dup.iloc[0]!r} is mapped to more than one field")
    return df


BROKEN_COLUMNS = ["field", "caption", "sheets", "dashboards", "calculations", "filters"]


def _usage_of(template: Template) -> dict:
    """`field_usage` of the template's workbook by (datasource, field), and which worksheets filter on
    each field. Worked out once per template."""
    cached = template._cache.get("usage")
    if cached is None:
        usage = {(r.datasource, r.field): r for r in field_usage(template.parser).itertuples(index=False)}
        filters: dict[tuple, set] = {}
        for ws in template.parser.xml_doc.xpath("/workbook/worksheets/worksheet[@name]"):
            for flt in ws.xpath(".//filter[@column]"):
                m = re.match(r"\[([^\]]+)\]\.\[([^\]]+)\]", flt.get("column"))
                if not m:
                    continue
                ds, inst = m.groups()
                base = ws.xpath(".//datasource-dependencies[@datasource=$d]/column-instance[@name=$i]/@column",
                                d=ds, i=f"[{inst}]")
                filters.setdefault((ds, base[0] if base else f"[{inst}]"), set()).add(ws.get("name"))
        cached = template._cache["usage"] = (usage, filters)
    return cached


def broken_sheets(template: Template, mapping: pd.DataFrame, datasource: Optional[str] = None) -> pd.DataFrame:
    """Required fields with no column, and what will not work without them: the worksheets and dashboards
    that use the field, the calculations, sets and groups built on it, and the worksheets that filter on it.
    Parameters are not followed. Returns `field, caption, sheets, dashboards, calculations, filters`, the
    last four as `; `-joined names."""
    entry = template.datasource(datasource)
    usage, filters = _usage_of(template)
    mapped = {r["field"] for r in mapping.to_dict("records") if r.get("mapped_to")}
    rows = []
    for f in entry["fields"]:
        if not f["required"] or f["name"] in mapped:
            continue
        u = usage.get((entry["name"], f["name"]))
        rows.append({"field": f["name"], "caption": f.get("caption"), "sheets": "; ".join(f.get("used_by") or []),
                     "dashboards": "; ".join(u.dashboards) if u else "",
                     "calculations": "; ".join(u.calculations) if u else "",
                     "filters": "; ".join(sorted(filters.get((entry["name"], f["name"]), ())))})
    return pd.DataFrame(rows, columns=BROKEN_COLUMNS)


# ------------------------------------------------------- checks and explain --

CHECK_COLUMNS = ["check", "severity", "field", "column", "detail"]
EXPLAIN_COLUMNS = ["change", "severity", "kind", "object", "detail"]
_SEVERITY = {"error": 0, "warning": 1, "info": 2}
_KEY_WORDS = {"id", "key", "code", "uuid", "guid", "number", "no", "num"}
_SAMPLE_ROWS = 2000


def _label(f: dict) -> str:
    return f.get("caption") or f["name"].strip("[]")


def _chosen(mapping: pd.DataFrame, data: DataSource) -> dict[str, str]:
    return {r["field"]: data.column(r["mapped_to"])
            for r in load_mapping(mapping).to_dict("records") if r.get("mapped_to")}


def check_data(
    template: Template,
    data: DataSource,
    mapping: pd.DataFrame,
    datasource: Optional[str] = None,
    deep: bool = False,
) -> pd.DataFrame:
    """What about the new data may make the workbook misleading even though it opens: findings only, never
    a reason to refuse an apply. Returns `check, severity, field, column, detail`.

    `dimension-missing`: a dimension the template's sheets use has no column, so the data may be at another
    grain. For a CSV also `empty-column` (a required field's column has no values) and `duplicate-key`
    (a column that looks like a key, `*_id`, `*_key`, `*_code`..., repeats a value). A CSV is judged on its
    first 2000 rows unless `deep`, which reads all of it."""
    entry = template.datasource(datasource)
    chosen = _chosen(mapping, data)
    by_field = {f["name"]: f for f in entry["fields"]}
    rows = []
    for f in entry["fields"]:
        if f["required"] and f.get("role") == "dimension" and f["name"] not in chosen:
            rows.append({"check": "dimension-missing", "severity": "warning", "field": f["name"], "column": "",
                         "detail": f"{', '.join(f.get('used_by') or [])} group or filter by {_label(f)}, "
                                   "which the data has no column for; its rows may be at a different grain"})
    if data.kind == "csv":
        sep = "\t" if Path(data.path).suffix.lower() == ".tsv" else ","
        frame = pd.read_csv(data.path, sep=sep, dtype=str, keep_default_na=False, encoding="utf-8-sig",
                            nrows=None if deep else _SAMPLE_ROWS)
        scope = f"in all {len(frame)} rows" if deep or len(frame) < _SAMPLE_ROWS else f"in the first {_SAMPLE_ROWS} rows"
        for field, column in chosen.items():
            values = frame[column].str.strip()
            filled = values[values != ""]
            if by_field[field]["required"] and filled.empty:
                rows.append({"check": "empty-column", "severity": "warning", "field": field, "column": column,
                             "detail": f"{column!r} has no values {scope}, and {', '.join(by_field[field].get('used_by') or [])} use it"})
            words = _words(column)
            if words and words[-1].lower() in _KEY_WORDS and not filled.empty and filled.duplicated().any():
                repeated = int(filled.duplicated(keep=False).sum())
                rows.append({"check": "duplicate-key", "severity": "info", "field": field, "column": column,
                             "detail": f"{column!r} looks like a key but {repeated} of {len(filled)} values repeat {scope}"})
    return pd.DataFrame(rows, columns=CHECK_COLUMNS).sort_values(
        ["severity", "check", "field"], key=lambda s: s.map(_SEVERITY) if s.name == "severity" else s,
        kind="stable", ignore_index=True)


def explain(template: Template, data: DataSource, mapping: pd.DataFrame,
            datasource: Optional[str] = None) -> pd.DataFrame:
    """What an apply changes besides the connection, so nothing is a surprise. Returns `change,
    severity, kind, object, detail` (errors first). Not covered: parameter values (you pass them, and the
    answers record them; no parameter of any corpus workbook depends on a field), and the extracts and packaged
    data that `make_template` already left out.

    `type-changed`: a field takes the type of the column that now feeds it. `field-dropped`: a field has no
    column. `affected`: what that breaks, one row per worksheet, dashboard, calculation, set or group and
    filter that uses the dropped field, directly or through others. `mapping-note`: a field mapped by a close
    match or to a column of another role."""
    entry = template.datasource(datasource)
    mapping = load_mapping(mapping)
    chosen = _chosen(mapping, data)
    status = {r["field"]: r.get("status", "") for r in mapping.to_dict("records")}
    usage, filters = _usage_of(template)
    kind_of = {(u.datasource, u.caption or u.field.strip("[]")): u.kind for u in usage.values()}
    rows = []
    unused_dropped: list[str] = []

    def add(change, severity, kind, obj, detail):
        rows.append({"change": change, "severity": severity, "kind": kind, "object": obj, "detail": detail})

    for f in entry["fields"]:
        label = _label(f)
        if f["name"] in chosen:
            column = chosen[f["name"]]
            new = data.datatype(column)
            if new and new != f["datatype"] and not f.get("customized"):
                bad = _type_status(f.get("physical_type") or f["datatype"], new) == "type mismatch"
                add("type-changed", "warning" if bad else "info", "field", label,
                    f"{f['datatype']} -> {new}, from column {column!r}")
            if status.get(f["name"]) in ("close match", "role differs"):
                add("mapping-note", "info" if status[f["name"]] == "close match" else "warning", "field", label,
                    f"{status[f['name']]}: column {column!r}")
            continue
        used = f.get("used_by") or []
        if not f["required"] and not used:
            unused_dropped.append(label)       # one row for all of these: they change nothing a sheet shows
            continue
        add("field-dropped", "error" if f["required"] else "info", "field", label,
            "no column in the data" + (f"; used by {', '.join(used)}" if used else "; no sheet uses it"))
        u = usage.get((entry["name"], f["name"]))
        for sheet in used:
            add("affected", "error", "sheet", sheet, f"uses {label}")
        for dash in (u.dashboards if u else []):
            add("affected", "warning", "dashboard", dash, f"shows a sheet that uses {label}")
        for calc in (u.calculations if u else []):
            kind = "calculation" if kind_of.get((entry["name"], calc), "calculated") == "calculated" else "set or group"
            add("affected", "error", kind, calc, f"is built on {label}")
        for sheet in sorted(filters.get((entry["name"], f["name"]), ())):
            add("affected", "error", "filter", sheet, f"filters on {label}")
    if unused_dropped:
        shown = ", ".join(unused_dropped[:8]) + (f", and {len(unused_dropped) - 8} more" if len(unused_dropped) > 8 else "")
        add("field-dropped", "info", "fields", f"{len(unused_dropped)} fields no sheet uses",
            f"no column in the data: {shown}")
    return pd.DataFrame(rows, columns=EXPLAIN_COLUMNS).sort_values(
        ["severity", "change", "kind", "object"], key=lambda s: s.map(_SEVERITY) if s.name == "severity" else s,
        kind="stable", ignore_index=True)


# ----------------------------------------------------------------- apply --

def _param_literal(datatype: str, raw: str) -> str:
    """Tableau's literal for a parameter value: `5`, `"East"`, `#2026-01-31#`."""
    raw = str(raw).strip()
    if datatype in _NUMERIC:
        try:
            num = float(raw)
        except ValueError:
            raise TemplateError(f"{raw!r} is not a number") from None
        if datatype == "integer":
            if num != int(num):
                raise TemplateError(f"{raw!r} is not a whole number")
            return str(int(num))
        return repr(num) if num != int(num) else f"{num:.1f}"
    if datatype == "boolean":
        if raw.lower() not in ("true", "false"):
            raise TemplateError(f"{raw!r} is not true or false")
        return raw.lower()
    if datatype in _DATES:
        try:
            when = pd.Timestamp(raw)
        except (ValueError, TypeError):
            raise TemplateError(f"{raw!r} is not a date") from None
        return "#" + (when.strftime("%Y-%m-%d") if datatype == "date" else when.strftime("%Y-%m-%d %H:%M:%S")) + "#"
    return '"' + raw.replace('"', '""') + '"'


def _set_parameters(doc, params: dict[str, str], template: Template) -> dict[str, str]:
    known = {p["caption"]: p for p in template.manifest.get("parameters", [])}
    known.update({p["name"].strip("[]"): p for p in template.manifest.get("parameters", [])})
    applied = {}
    for key, raw in params.items():
        p = known.get(key) or known.get(key.strip("[]"))
        if p is None:
            raise TemplateError(f"template has no parameter {key!r}; it has: "
                                + ", ".join(q["caption"] for q in template.manifest.get("parameters", [])))
        try:
            literal = _param_literal(p["datatype"], raw)
        except TemplateError as e:
            raise TemplateError(f"parameter {p['caption']!r}: {e}") from None
        if p.get("allowed") and literal not in p["allowed"]:
            raise TemplateError(f"parameter {p['caption']!r}: {raw!r} is not one of its allowed values")
        for col in doc.xpath("/workbook/datasources/datasource[@name='Parameters']/column[@name=$n]", n=p["name"]):
            col.set("value", literal)
            calc = col.find("calculation")
            if calc is not None:
                calc.set("formula", literal)
        applied[p["caption"]] = literal
    return applied


def _connection_id(kind: str, path: str) -> str:
    return f"{kind}." + hashlib.sha1(path.encode("utf-8")).hexdigest()[:28]


_OM = "_.fcp.ObjectModelEncapsulateLegacy"   # Tableau 2020.2+ object model feature flags
_OM_TABLE = "_.fcp.ObjectModelTableType.true...column"


def _object_model(doc) -> Optional[str]:
    """How a workbook writes the 2020.2+ object model: `"prefixed"` (feature-flag tags such as
    `_.fcp.ObjectModelEncapsulateLegacy.true...object-graph`), `"plain"` (an ordinary
    `<object-graph>`; same format version, newer builds) or None for neither. Both occur in the
    corpus, and a workbook's sheets reference the table column either way."""
    if any(isinstance(el.tag, str) and el.tag.startswith(_OM) for el in doc.iter()):
        return "prefixed"
    if doc.xpath("/workbook/datasources/datasource[object-graph or column[@datatype='table']]"):
        return "plain"
    return None


def _csv_connection(data: DataSource, local_of: dict[str, str], model: Optional[str] = "prefixed",
                    object_id: Optional[str] = None):
    """A federated connection to a CSV file, in the shape Tableau writes.

    With an object `model` (see `_object_model`: the template uses Tableau's
    2020.2+ object model) it also returns the table column and object graph
    the datasource needs; the prefixed form writes the legacy and
    object-model relations side by side, as Tableau does, the plain form
    one relation."""
    p = Path(data.path)
    # Worksheets name the table column (record counts) by this id, so a template that has one keeps it.
    object_id = object_id or (f"{re.sub(r'[^0-9A-Za-z_]', '_', p.stem)}_"
                              + hashlib.md5(data.path.encode("utf-8")).hexdigest().upper())
    return _file_connection(
        data, local_of, model, object_id, caption=p.stem, table=p.name,
        conn_id=_connection_id("textscan", data.path),
        named={"class": "textscan", "directory": p.parent.as_posix(), "filename": p.name, "password": "", "server": ""},
        rel_table=f"[{p.stem}#{p.suffix.lstrip('.')}]",
        cols={"character-set": "UTF-8", "header": "yes", "locale": "en_US",
              "separator": "\t" if p.suffix.lower() == ".tsv" else ","},
        record=lambda f: (str(_REMOTE_TYPE.get(f["datatype"], 129)), "Sum" if f["datatype"] in _NUMERIC else "Count", None),
    )


def _excel_connection(data: DataSource, local_of: dict[str, str], model: Optional[str] = "prefixed",
                      object_id: Optional[str] = None):
    """A federated connection to one worksheet of an Excel file (`excel-direct`), in the shape Tableau writes
    it (copied from the 88 `excel-direct` workbooks of the corpus: the relation is `[Sheet$]`, its `columns`
    carry the sheet's `gridOrigin`, the remote types are the Excel driver's). Same object-model forms as
    `_csv_connection`."""
    p = Path(data.path)
    sheet = data.sheet or ""
    object_id = object_id or (f"{re.sub(r'[^0-9A-Za-z_]', '_', sheet)}_"
                              + hashlib.md5(f"{data.path}\x1f{sheet}".encode("utf-8")).hexdigest().upper())
    grid = data.grid or "A1:A1"

    def record(f):
        code, debug = _EXCEL_REMOTE_TYPE.get(f["datatype"], (130, "WSTR"))
        aggregation = "Sum" if f["datatype"] in _NUMERIC else "Year" if f["datatype"] in _DATES else "Count"
        return str(code), aggregation, debug

    return _file_connection(
        data, local_of, model, object_id, caption=sheet or p.stem, table=sheet,
        conn_id=_connection_id("excel-direct", data.path),
        named={"class": "excel-direct", "cleaning": "no", "compat": "no", "dataRefreshTime": "", "filename": p.as_posix(),
               "interpretationMode": "0", "password": "", "server": "", "validate": "no"},
        named_caption=p.stem, rel_table=f"[{sheet}$]",
        cols={"gridOrigin": f"{grid}:no:{grid}:0", "header": "yes", "outcome": "2"}, record=record,
    )


def _file_connection(data: DataSource, local_of: dict[str, str], model: Optional[str], object_id: str, caption: str,
                     table: str, conn_id: str, named: dict, rel_table: str, cols: dict, record,
                     named_caption: Optional[str] = None):
    """The part of a file connection `_csv_connection` and `_excel_connection` share: the named connection, the
    relation (twice for the prefixed object model), one metadata record per column and, for an object model, the
    table column and object graph. `record(field)` gives the remote type, aggregation and, if the driver writes
    one, the `DebugRemoteType` text."""
    modern = model is not None
    om = (lambda name: f"{_OM}.true...{name}") if model == "prefixed" else (lambda name: name)
    conn = etree.Element("connection", {"class": "federated"})
    node = etree.SubElement(etree.SubElement(conn, "named-connections"), "named-connection",
                            caption=named_caption or caption, name=conn_id)
    etree.SubElement(node, "connection", named)

    def relation():
        rel = etree.Element("relation", connection=conn_id, name=table, table=rel_table, type="table")
        columns = etree.SubElement(rel, "columns", cols)
        for i, f in enumerate(data.fields):
            etree.SubElement(columns, "column", datatype=f["datatype"], name=f["name"], ordinal=str(i))
        return rel

    if model == "prefixed":
        for flag in ("false", "true"):
            rel = relation()
            rel.tag = f"{_OM}.{flag}...relation"
            conn.append(rel)
    else:
        conn.append(relation())
    records = etree.SubElement(conn, "metadata-records")
    for i, f in enumerate(data.fields):
        remote_type, aggregation, debug = record(f)
        rec = etree.SubElement(records, "metadata-record", {"class": "column"})
        for tag, text in (
            ("remote-name", f["name"]), ("remote-type", remote_type),
            ("local-name", local_of[f["name"]]), ("parent-name", f"[{table}]"),
            ("remote-alias", f["name"]), ("ordinal", str(i)), ("local-type", f["datatype"]),
            ("aggregation", aggregation), ("contains-null", "true"),
        ):
            etree.SubElement(rec, tag).text = text
        if debug:
            attrs = etree.SubElement(rec, "attributes")
            etree.SubElement(attrs, "attribute", datatype="string", name="DebugRemoteType").text = f'"{debug}"'
        if modern:
            etree.SubElement(rec, om("object-id")).text = f"[{object_id}]"
    if not modern:
        return conn, []
    table_col = etree.Element(_OM_TABLE if model == "prefixed" else "column", caption=caption, datatype="table",
                              name=f"[__tableau_internal_object_id__].[{object_id}]",
                              role="measure", type="quantitative")
    graph = etree.Element(om("object-graph"))
    obj = etree.SubElement(etree.SubElement(graph, "objects"), "object", caption=caption, id=object_id)
    props = etree.SubElement(obj, "properties", context="")
    props.append(relation())
    return conn, [table_col, graph]


def _tableau_connection(data: DataSource, local_of: dict[str, str]):
    """The new data's own connection, with each column's local name set to
    the template field it now feeds."""
    src = data.element
    conn = copy.deepcopy(src.find("connection"))
    rename = {f["local"]: local_of[f["name"]] for f in data.fields}
    for rec in conn.xpath(".//metadata-record[@class='column']"):
        el = rec.find("local-name")
        if el is not None and el.text in rename:
            el.text = rename[el.text]
    for m in conn.xpath(".//cols/map[@key]"):
        if m.get("key") in rename:
            m.set("key", rename[m.get("key")])
    # The object model (2020.2+) describes the tables of this connection.
    extras = [copy.deepcopy(el) for el in src if _endswith_tag(el, "object-graph")
              or (isinstance(el.tag, str) and el.tag.endswith("column") and el.get("datatype") == "table")]
    return conn, extras


# --------------------------------------------------------------- answers --

ANSWERS_FORMAT = TEMPLATE_FORMAT + "-answers"
ANSWERS_VERSION = 2   # 2 adds the template id/revision, per-datasource entries, profiles and a schema fingerprint
_CREDENTIAL_KEYS = ("password", "username", "token", "secret")
# Keys under these hold names the user chose (parameter captions, field names): not part of the format.
_FREE_KEYS = ("parameters", "mapping")


def _reject_credentials(node, path: str = "") -> None:
    """An answers file is shared and kept in version control: it never holds credentials. Credentials are
    entered in Tableau on opening the workbook, as with a Power BI template."""
    if isinstance(node, dict):
        for key, value in node.items():
            if str(key).lower() in _CREDENTIAL_KEYS:
                raise TemplateError(f"answers must not hold credentials, but have {path}{key!r}; "
                                    "enter them in Tableau when you open the workbook")
            if key not in _FREE_KEYS:
                _reject_credentials(value, f"{path}{key}.")
    elif isinstance(node, list):
        for i, item in enumerate(node):
            _reject_credentials(item, f"{path}{i}.")


def load_answers(source: Union[str, os.PathLike, dict]) -> dict:
    """Answers for `apply_template`: a dict, a standalone `*.answers.json`, or a workbook made by
    `apply_template` (its `template-answers.json`). Version 1 answers load too."""
    if isinstance(source, dict):
        answers = copy.deepcopy(source)
    else:
        path = Path(source)
        if path.suffix.lower() == ".twbx":
            answers = read_answers(str(path))
            if answers is None:
                raise TemplateError(f"{path.name} has no {ANSWERS_NAME}; it was not made by `template apply`")
        else:
            try:
                answers = json.loads(path.read_text(encoding="utf-8-sig"))
            except FileNotFoundError:
                raise
            except json.JSONDecodeError as e:
                raise TemplateError(f"{path.name} is not valid JSON: {e}") from None
        if isinstance(answers, dict):
            answers.setdefault("_base_dir", str(path.resolve().parent))
    if not isinstance(answers, dict) or answers.get("format") != ANSWERS_FORMAT:
        raise TemplateError(f"not an answers file (format {ANSWERS_FORMAT!r} expected)")
    if int(answers.get("version", 0)) > ANSWERS_VERSION:
        raise TemplateError(f"answers were made by a newer py-tbparse (answers v{answers['version']})")
    _reject_credentials(answers)
    return answers


def _answers_entries(answers: dict) -> list[dict]:
    """The per-datasource entries (`datasource`, `data`, `mapping`, `missing`) of answers. Version 1
    answers held one datasource at the top level."""
    if answers.get("datasources"):
        return [copy.deepcopy(e) for e in answers["datasources"]]
    if answers.get("datasource"):
        return [{"datasource": answers["datasource"], "data": answers.get("data") or {},
                 "mapping": answers.get("mapping") or {}, "missing": answers.get("missing") or []}]
    return []


def _answers_entry(answers: dict, datasource: str) -> dict:
    """What the answers say about one template datasource, or {}."""
    return next((e for e in _answers_entries(answers) if e.get("datasource") == datasource), {})


def _raw_value(datatype: str, literal: str) -> str:
    """The inverse of `_param_literal`: `"East"` -> `East`, `#2026-01-31#` -> `2026-01-31`."""
    literal = str(literal)
    if datatype in _DATES and literal.startswith("#") and literal.endswith("#"):
        return literal[1:-1]
    if datatype not in _NUMERIC and datatype != "boolean" and literal.startswith('"') and literal.endswith('"'):
        return literal[1:-1].replace('""', '"')
    return literal


def _answers_parameters(answers: dict, template: Template) -> dict[str, str]:
    """The parameter values saved in answers (stored as Tableau literals), as the values a caller passes."""
    types = {p["caption"]: p["datatype"] for p in template.manifest.get("parameters", [])}
    return {k: _raw_value(types.get(k, "string"), v) for k, v in (answers.get("parameters") or {}).items()}


def _overlay_mapping(mapping: pd.DataFrame, saved: dict, data: DataSource, entry: dict) -> tuple[pd.DataFrame, list]:
    """The suggested mapping with the saved choices put back. A saved column the data no longer has, or a
    field the template no longer has, is left to the suggestion and reported as stale."""
    mapping = mapping.copy()
    fields = {f["name"] for f in entry["fields"]}
    columns = set(data.names())
    stale = []
    for field, column in saved.items():
        if field not in fields or column not in columns:
            stale.append(field)
            continue
        for i in mapping.index[(mapping["mapped_to"] == column) & (mapping["field"] != field)]:
            mapping.loc[i, ["mapped_to", "data_type", "score"]] = ["", "", None]    # the column now feeds another field
            mapping.loc[i, "status"] = "missing" if mapping.loc[i, "required"] else "unused"
        row = mapping["field"] == field
        mapping.loc[row, ["mapped_to", "data_type", "status", "score"]] = [column, data.datatype(column), "from answers", 1.0]
    return mapping, sorted(stale)


def default_output_path(template: Template, data: DataSource) -> str:
    stem = Path(template.path).name
    for suffix in (".template.twbx", ".twbx"):
        if stem.lower().endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    return str(Path(template.path).with_name(f"{stem}_{Path(data.path).stem}.twbx"))


@dataclass
class ApplyPlan:
    """Everything `apply_template` has settled before it writes: the template and its datasource entry, the
    data, the mapping (suggested, edited or from answers) and the parameter values to set."""

    template: Template
    entry: dict
    data: DataSource
    mapping: pd.DataFrame
    params: dict
    saved: Optional[dict] = None       # the answers, if any
    stale: list = _field(default_factory=list)    # saved mapping entries that no longer fit
    changed: bool = False              # the data's columns differ from the answers' last run


def resolve_apply(
    template: Union[Template, str],
    data: Union[DataSource, str, None] = None,
    mapping: Optional[pd.DataFrame] = None,
    params: Optional[dict[str, str]] = None,
    datasource: Optional[str] = None,
    data_datasource: Optional[str] = None,
    answers: Union[str, os.PathLike, dict, None] = None,
    profile: Optional[str] = None,
    fuzzy_cutoff: float = 0.85,
    sheet: Union[str, int, None] = None,
) -> ApplyPlan:
    """What `apply_template` would do, without writing: the same arguments, the same precedence (explicit
    argument, then profile, then saved answers), the same errors."""
    if not isinstance(template, Template):
        template = load_template(template)
    saved = load_answers(answers) if answers is not None else None
    chosen_profile: dict = {}
    if profile is not None:
        if saved is None:
            raise TemplateError("a profile is a named set inside an answers file; pass answers= too")
        profiles = saved.get("profiles") or {}
        if profile not in profiles:
            raise TemplateError(f"answers have no profile {profile!r}; they have: "
                                + (", ".join(sorted(profiles)) or "none"))
        chosen_profile = profiles[profile]
    if saved is not None and template.id and saved.get("template", {}).get("id") not in (None, template.id):
        raise TemplateError("these answers were made for another template "
                            f"(id {saved['template']['id']}, this one is {template.id})")
    if datasource is None and saved is not None and saved.get("datasource"):
        datasource = saved["datasource"]
    entry = template.datasource(datasource)
    prior = _answers_entry(saved, entry["name"]) if saved is not None else {}
    if data is None:
        file = chosen_profile.get("data") or (prior.get("data") or {}).get("file")
        if not file:
            raise TemplateError("no data: pass data=, or answers/profile that name a data file")
        if not Path(file).is_absolute() and saved and saved.get("_base_dir"):
            file = str(Path(saved["_base_dir"]) / file)
        if data_datasource is None and not chosen_profile.get("data"):
            data_datasource = (prior.get("data") or {}).get("datasource")
        if sheet is None and not chosen_profile.get("data"):
            sheet = (prior.get("data") or {}).get("sheet")
        data = file
    if not isinstance(data, DataSource):
        data = read_data(data, datasource=data_datasource, sheet=sheet)
    stale: list = []
    if mapping is None:
        mapping = suggest_mapping(template, data, datasource=entry["name"], fuzzy_cutoff=fuzzy_cutoff)
        if prior.get("mapping"):
            mapping, stale = _overlay_mapping(mapping, prior["mapping"], data, entry)
    else:
        mapping = load_mapping(mapping)
    changed = bool(prior.get("data", {}).get("schema_fingerprint")) and \
        prior["data"]["schema_fingerprint"] != data.fingerprint()
    merged = {**(_answers_parameters(saved, template) if saved is not None else {}),
              **(chosen_profile.get("parameters") or {}), **(params or {})}
    return ApplyPlan(template=template, entry=entry, data=data, mapping=mapping, params=merged,
                     saved=saved, stale=stale, changed=changed)


def apply_template(
    template: Union[Template, str],
    data: Union[DataSource, str, None] = None,
    mapping: Optional[pd.DataFrame] = None,
    params: Optional[dict[str, str]] = None,
    output_path: Optional[str] = None,
    datasource: Optional[str] = None,
    data_datasource: Optional[str] = None,
    allow_missing: bool = False,
    overwrite: bool = False,
    report: Optional[dict] = None,
    answers: Union[str, os.PathLike, dict, None] = None,
    profile: Optional[str] = None,
    sheet: Union[str, int, None] = None,
) -> str:
    """Make a new workbook from a template and new data; return its path.

    `mapping` defaults to `suggest_mapping(template, data)`; pass an edited
    one (see `load_mapping`) to choose columns yourself. A required field
    with no column would break the sheets listed by `broken_sheets`, so that
    raises `TemplateError` unless `allow_missing=True`. `params` sets
    parameter values by caption (`{"Top N": "10"}`), checked against the
    parameter's type and allowed values.

    `answers` runs this without asking: a `*.answers.json`, a workbook made
    by an earlier `apply_template`, or a dict (see `load_answers`). Its saved
    mapping is put over the suggestion (a saved column the data no longer
    has is reported as `report["stale_mapping"]`, and `report["schema_changed"]`
    says whether the data's columns differ from last time), its parameters
    are set, and `data` may be left out to use the file it names. `profile`
    picks a named set from the answers' `profiles` (`{"prod": {"parameters":
    {...}, "data": "prod.csv"}}`). An explicit argument beats the profile,
    which beats the saved answers. Answers never hold credentials.

    The template datasource's connection is replaced by one to the new data
    (a CSV file, one worksheet of an Excel file -- `sheet=`, see `read_data` --
    or the connection of the given workbook / .tds); every
    field keeps the local name its sheets and formulas use. The answers
    (template, data, mapping, parameters) are saved inside the output as
    `template-answers.json`. Output is a `.twbx` (default
    `<template>_<data>.twbx` beside the template) and is never overwritten
    unless `overwrite=True`.
    """
    plan = resolve_apply(template, data, mapping=mapping, params=params, datasource=datasource,
                         data_datasource=data_datasource, answers=answers, profile=profile, sheet=sheet)
    template, data, entry, mapping, params = plan.template, plan.data, plan.entry, plan.mapping, plan.params
    saved, stale, changed = plan.saved, plan.stale, plan.changed
    by_field = {f["name"]: f for f in entry["fields"]}
    chosen: dict[str, str] = {}
    for r in mapping.to_dict("records"):
        if r["field"] not in by_field:
            raise TemplateError(f"mapping names {r['field']!r}, which the template does not have")
        if r.get("mapped_to"):
            column = data.column(r["mapped_to"])
            if column not in data.names():
                raise TemplateError(f"mapping uses column {r['mapped_to']!r}, which {Path(data.path).name} does not have")
            chosen[r["field"]] = column
    missing = [f for f in entry["fields"] if f["required"] and f["name"] not in chosen]
    if missing and not allow_missing:
        sheets = sorted({s for f in missing for s in f.get("used_by") or []})
        raise TemplateError(
            "no column for required field(s) "
            + ", ".join(f.get("caption") or f["name"].strip("[]") for f in missing)
            + f"; sheets that would break: {', '.join(sheets)} (map them, or pass allow_missing)"
        )

    # local name for every column of the new data
    field_of = {col: fld for fld, col in chosen.items()}
    reserved = {f["name"] for f in entry["fields"]}
    local_of: dict[str, str] = {}
    for f in data.fields:
        if f["name"] in field_of:
            local_of[f["name"]] = field_of[f["name"]]
            continue
        local = "[" + f["name"].replace("]", "]]") + "]"
        while local in reserved or local in local_of.values():
            local = local[:-1] + " (data)]"
        local_of[f["name"]] = local

    doc = copy.deepcopy(template.parser.xml_doc)
    ds_el = doc.xpath("/workbook/datasources/datasource[@name=$n]", n=entry["name"])[0]
    if data.kind in ("csv", "excel"):
        # a column with no values takes the type of the field it feeds (else string)
        typed = copy.copy(data)
        typed.fields = [{**f, "datatype": f["datatype"] or (
            by_field[field_of[f["name"]]].get("physical_type") or by_field[field_of[f["name"]]]["datatype"]
            if f["name"] in field_of else "string")} for f in data.fields]
        old_ids = ds_el.xpath("./*[substring(name(), string-length(name()) - 11) = 'object-graph']"
                              "/objects/object/@id")
        build = _excel_connection if data.kind == "excel" else _csv_connection
        conn, extras = build(typed, local_of, model=_object_model(doc), object_id=old_ids[0] if old_ids else None)
    else:
        conn, extras = _tableau_connection(data, local_of)
    old = ds_el.find("connection")
    for el in list(ds_el):
        if el is not old and (_endswith_tag(el, "extract") or _endswith_tag(el, "object-graph")
                              or (isinstance(el.tag, str) and el.tag.endswith("column")
                                  and el.get("datatype") == "table")):
            ds_el.remove(el)
    if old is not None:
        old.addprevious(conn)
        ds_el.remove(old)
    else:
        ds_el.insert(0, conn)
    for el in extras:
        if _endswith_tag(el, "object-graph"):
            ds_el.append(el)  # Tableau writes it last
            continue
        cols = [c for c in ds_el if isinstance(c.tag, str) and (c.tag == "column" or c.tag.endswith("...column"))]
        # Tableau puts a table column after <aliases> (all 191 object-model datasources of the corpus)
        aliases = ds_el.find("aliases")
        (cols[-1].addnext(el) if cols else (conn if aliases is None else aliases).addnext(el))
    # a field now fed by a column of another type takes that type, unless the
    # author set the field's type by hand (Tableau then converts the column)
    for fld, col in chosen.items():
        new_type = data.datatype(col)
        for c in ds_el.xpath("./column[@name=$n]", n=fld):
            if new_type and c.get("datatype") != new_type and c.get("datatype-customized") != "true":
                c.set("datatype", new_type)
    # Tableau's own record count is a formula (`1`) with no data column, as in all 42 built-ins of the
    # corpus. A template leaves out a data column that shared its name, so without the formula a worksheet
    # that sums it has nothing to sum.
    fed = {r.findtext("local-name") for r in ds_el.xpath("./connection//metadata-record[@class='column']")}
    for c in ds_el.xpath("./column[@name]"):
        if c.get(_AUTO_COLUMN) == "numrec" and c.find("calculation") is None and c.get("name") not in fed:
            etree.SubElement(c, "calculation", {"class": "tableau", "formula": "1"})

    applied_params = _set_parameters(doc, params, template)
    this = {
        "datasource": entry["name"],
        "data": {"file": data.path, "kind": data.kind,
                 "datasource": data.element.get("name") if data.element is not None else None,
                 "schema_fingerprint": data.fingerprint(), "columns": data.names(),
                 **({"sheet": data.sheet} if data.sheet else {})},
        "mapping": {fld: col for fld, col in chosen.items()},
        "missing": [f["name"] for f in missing],
    }
    order = [e["name"] for e in template.manifest.get("datasources", [])]
    others = [e for e in _answers_entries(saved) if e.get("datasource") != entry["name"]] if saved is not None else []
    entries = sorted(others + [this], key=lambda e: order.index(e["datasource"]) if e["datasource"] in order else len(order))
    out_answers = {
        "format": ANSWERS_FORMAT,
        "version": ANSWERS_VERSION,
        "created": _now(),
        "created_with": f"py-tbparse {_version()}",
        "template": {"id": template.id, "revision": template.revision, "name": template.name,
                     "file": Path(template.path).name, "manifest_sha256": template.manifest_sha256,
                     # what the template needed then, so `template update` can say what changed since
                     "manifest": template.manifest},
        # the datasource this run filled, as version 1 kept it; `datasources` has them all
        "data": this["data"], "datasource": this["datasource"], "mapping": this["mapping"], "missing": this["missing"],
        "parameters": applied_params,
        "datasources": entries,
        "profiles": copy.deepcopy((saved or {}).get("profiles") or {}),
    }
    if profile is not None:
        out_answers["profile"] = profile
    _reject_credentials(out_answers)
    twb = etree.tostring(doc, xml_declaration=True, encoding="utf-8")
    out = Path(output_path) if output_path else Path(default_output_path(template, data))
    if out.suffix.lower() != ".twbx":
        raise TemplateError(f"output must be a .twbx file, got {out.suffix or 'no extension'}")
    if out.resolve() in (Path(template.path).resolve(), Path(data.path).resolve()):
        raise FileExistsError(f"refusing to overwrite an input: {out}")
    packed = _package(template.parser, twb, {ANSWERS_NAME: json.dumps(out_answers, indent=2).encode("utf-8")},
                      keep_data=False, drop=[MANIFEST_NAME])
    _write_new(out, packed, overwrite)
    if report is not None:
        report.update(mapped=len(chosen), missing=len(missing), parameters=len(applied_params))
        if saved is not None:
            report.update(stale_mapping=stale, schema_changed=changed)
    return str(out)


def read_answers(path: str) -> Optional[dict]:
    """The `template-answers.json` of a workbook made by `apply_template`, or None."""
    with zipfile.ZipFile(path) as z:
        if ANSWERS_NAME not in z.namelist():
            return None
        return json.loads(z.read(ANSWERS_NAME).decode("utf-8"))
